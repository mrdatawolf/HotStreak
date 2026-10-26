from datetime import date

import openpyxl
from openpyxl.styles import PatternFill
from sqlalchemy import select

from app import models
from app.sheet_import import check_totals, parse_tabs, read_workbook, split_client, split_items, sync

TODAY = date(2026, 10, 23)


def save(tmp_path, **tabs):
    """Write an .xlsx shaped like the client's and read it back.

    Each tab is {"B2": value} or {"B2": (value, fill)}; pass {"quota_col": "P"} to move QUOTA/SALES.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        rows = dict(rows)
        quota_col = rows.pop("quota_col", "N")
        ws["B1"], ws[f"{quota_col}1"] = "MONTH", "QUOTA"
        ws.cell(1, ws[f"{quota_col}1"].column + 1, "SALES")
        for ref, value in rows.items():
            if isinstance(value, tuple):
                value, fill = value
                ws[ref].fill = PatternFill("solid", fgColor=fill)
            ws[ref] = value
    path = tmp_path / "sheet.xlsx"
    wb.save(path)
    return read_workbook(path)


SEPT = {
    "A2": "KANDIS", "B2": "Pacific Union School", "B3": "(2) BP-70M75", "B4": 27968.10,
    "C2": "Aquaterra*", "C3": "Stand", "C4": 442.0, "N2": 44, "O2": 28410.10,
    "A5": "HOUSE", "B5": "Glenn County+", "B6": "Various", "B7": 1000.0, "N5": "~", "O5": 1000.0,
    "A8": "MONTHLY TOTAL", "B8": 29410.10,
}


def test_split_items():
    assert split_items("(2) BP-71M45  XM3250") == [(2, "BP-71M45"), (1, "XM3250")]
    assert split_items("Lexmark M3250 (2)") == [(2, "Lexmark M3250")]
    assert split_items("(2) XC4342 & Etherfax") == [(2, "XC4342"), (1, "Etherfax")]
    assert split_items("BP-70C45 (8) Lexmark XM3250\nMX-C407F") == [(1, "BP-70C45"), (8, "Lexmark XM3250"), (1, "MX-C407F")]
    assert split_items("(1)    MX-1056") == [(1, "MX-1056")]
    assert split_items('48 Mo. Rental Sharp C357F') == [(1, "48 Mo. Rental Sharp C357F")]
    assert split_items(None) == []


def test_split_client():
    assert split_client("Aquaterra*") == ("Aquaterra", True, False)
    assert split_client("RRHC+ ") == ("RRHC", False, True)
    assert split_client("  Bear  River ") == ("Bear River", False, False)


def test_parse_month(tmp_path):
    result = parse_tabs(save(tmp_path, **{"926": SEPT, "Overview": {"A1": "x"}}))
    assert result.warnings == []
    [month] = result.months
    assert (month.year, month.month, month.sheet_total) == (2026, 9, 29410.10)
    kandis, house = month.reps
    assert (kandis.name, kandis.quota, house.quota) == ("KANDIS", 44, None)
    school, aquaterra = kandis.deals
    assert (school.client, school.items, school.amount, school.where) == ("Pacific Union School", [(2, "BP-70M75")], 27968.10, "926!B2")
    assert aquaterra.new_customer and not aquaterra.plus
    assert house.deals[0].plus
    assert check_totals(result) == []


def test_finds_moved_quota_column_and_blue_new_customer(tmp_path):
    feb = {"quota_col": "P", "A2": "DAVE", "B2": ("Corning High", "FFA4C2F4"), "B4": 100.0, "P2": 40,
           "A5": "MONTHLY TOTAL", "B5": 100.0}
    [month] = parse_tabs(save(tmp_path, **{"0224": feb})).months
    assert (month.year, month.month) == (2024, 2)
    [deal] = month.reps[0].deals
    assert deal.new_customer and month.reps[0].quota == 40


def test_check_totals_reports_mismatch(tmp_path):
    result = parse_tabs(save(tmp_path, **{"926": {**SEPT, "O2": 1.0}}))
    [problem] = check_totals(result)
    assert "KANDIS" in problem and "1.00" in problem


AUG = {"A2": "Kandis", "B2": "PACIFIC UNION SCHOOL", "B4": 10.0, "N2": 40, "A5": "MONTHLY TOTAL"}


def test_sync(tmp_path, session):
    result = parse_tabs(save(tmp_path, **{"926": SEPT, "826": AUG}))
    counts = sync(session, result, TODAY)
    session.commit()
    assert counts == {"months": 2, "reps": 2, "clients": 3, "deals": 4}

    periods = session.scalars(select(models.Period).order_by(models.Period.month)).all()
    assert [p.locked_at is not None for p in periods] == [True, False]
    house = session.scalar(select(models.Rep).where(models.Rep.name == "HOUSE"))
    assert house.counts_toward_total
    # Case-only differences are one client, named by the most-used spelling (tie → first seen).
    names = sorted(c.name for c in session.scalars(select(models.Client)))
    assert names == ["Aquaterra", "Glenn County", "PACIFIC UNION SCHOOL"]
    aquaterra = session.scalar(select(models.Deal).join(models.Client).where(models.Client.name == "Aquaterra"))
    assert [f.name for f in aquaterra.flags] == ["New customer"]
    # "+" is recorded on the deal, not as a flag, and stripped from the name.
    glenn = session.scalar(select(models.Deal).join(models.Client).where(models.Client.name == "Glenn County"))
    assert glenn.sheet_plus and glenn.flags == [] and not aquaterra.sheet_plus
    school = session.scalars(select(models.Deal).where(models.Deal.total_cents == 2796810)).one()
    assert [(line.qty, line.description) for line in school.lines] == [(2, "BP-70M75")]
    assert (school.sheet_ref, school.sheet_client) == ("926!B2", "Pacific Union School")


def test_sync_merges_repeated_rep(tmp_path, session):
    dec = {"A2": "RANEE", "B2": "A", "B4": 1.0, "N2": 40, "A5": "Ranee", "B5": "B", "B7": 2.0, "N5": 40,
           "A8": "MONTHLY TOTAL"}
    result = parse_tabs(save(tmp_path, **{"1225": dec}))
    sync(session, result, TODAY)
    assert any("RANEE appears twice" in w for w in result.warnings)
    [slot] = session.scalars(select(models.PeriodRep)).all()
    assert slot.quota == 40
    assert len(session.scalars(select(models.Deal)).all()) == 2


def test_resync_replaces_sheet_deals_and_keeps_app_deals(tmp_path, session):
    sync(session, parse_tabs(save(tmp_path, **{"926": SEPT})), TODAY)
    period = session.scalar(select(models.Period))
    kandis = session.scalar(select(models.Rep).where(models.Rep.name == "KANDIS"))
    app_deal = models.Deal(period=period, rep=kandis, client=models.Client(name="Walk-in"), total_cents=500)
    session.add(app_deal)
    session.commit()

    edited = {**SEPT, "C4": 500.0, "O2": 28468.10, "B8": 29468.10}
    counts = sync(session, parse_tabs(save(tmp_path, **{"926": edited})), TODAY)
    session.commit()
    assert counts["deals"] == 3
    deals = session.scalars(select(models.Deal)).all()
    assert len(deals) == 4 and app_deal in deals
    assert len(session.scalars(select(models.Period)).all()) == 1
    assert len(session.scalars(select(models.Rep)).all()) == 2
    assert len(session.scalars(select(models.PeriodRep)).all()) == 2
    aquaterra = session.scalar(select(models.Deal).where(models.Deal.sheet_ref == "926!C2"))
    assert aquaterra.total_cents == 50000


def test_sync_locks_all_but_current_month(tmp_path, session):
    # A tab set up early for next month doesn't take over the board.
    tabs = {"826": AUG, "926": SEPT, "1126": AUG}
    sync(session, parse_tabs(save(tmp_path, **tabs)), TODAY)
    locked = {p.month: p.locked_at is not None for p in session.scalars(select(models.Period))}
    assert locked == {8: True, 9: False, 11: True}
    # Next month: 11/26 opens, 9/26 locks.
    sync(session, parse_tabs(save(tmp_path, **tabs)), date(2026, 11, 2))
    locked = {p.month: p.locked_at is not None for p in session.scalars(select(models.Period))}
    assert locked == {8: True, 9: True, 11: False}
